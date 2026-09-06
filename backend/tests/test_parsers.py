from io import BytesIO

import docx
import openpyxl
from reportlab.pdfgen import canvas

from app.services import parsers
from app.services.contracts import Criterion
from app.services.llm import _normalize_criteria
from app.services.parsers import extract_url_text, fallback_rubric, parse_document


def test_parse_pdf_bytes_returns_text_and_metadata() -> None:
    buffer = BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(72, 720, "PDF assignment text")
    document.save()

    parsed = parse_document(buffer.getvalue(), "assignment.pdf")

    assert "PDF assignment text" in parsed.text
    assert parsed.metadata == {"format": "pdf", "pages": 1}


def test_parse_pdf_stream_rewinds_before_reading() -> None:
    buffer = BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(72, 720, "Rewound PDF assignment text")
    document.save()
    buffer.seek(0, 2)

    parsed = parse_document(buffer, "assignment.pdf")

    assert "Rewound PDF assignment text" in parsed.text


def test_parse_docx_bytes_returns_text_and_tables() -> None:
    buffer = BytesIO()
    document = docx.Document()
    document.add_paragraph("DOCX assignment text")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Criterion"
    table.cell(0, 1).text = "10"
    document.save(buffer)

    parsed = parse_document(buffer.getvalue(), "assignment.docx")

    assert "DOCX assignment text" in parsed.text
    assert "Criterion | 10" in parsed.text
    assert parsed.metadata["format"] == "docx"


def test_parse_docx_preserves_paragraph_breaks() -> None:
    buffer = BytesIO()
    document = docx.Document()
    document.add_paragraph("First requirement")
    document.add_paragraph("")
    document.add_paragraph("Second requirement")
    document.save(buffer)

    parsed = parse_document(buffer.getvalue(), "assignment.docx")

    assert "First requirement\n\nSecond requirement" in parsed.text


def test_parse_xlsx_bytes_returns_text_and_sheet_metadata() -> None:
    buffer = BytesIO()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Rubric"
    sheet.append(["Criterion", "Points"])
    sheet.append(["Analysis", 10])
    workbook.save(buffer)

    parsed = parse_document(buffer.getvalue(), "assignment.xlsx")

    assert "# Rubric" in parsed.text
    assert "Analysis | 10" in parsed.text
    assert parsed.metadata["sheets"] == ["Rubric"]


def test_parse_txt_removes_null_bytes_and_controls() -> None:
    parsed = parse_document(b"answer\x00 with\x01 controls\nsecond line", "answer.txt")

    assert parsed.text == "answer with controls\nsecond line"
    assert parsed.metadata == {"format": "txt"}


def test_fallback_rubric_infers_criteria_for_unstructured_task() -> None:
    rubric = fallback_rubric(
        "preview",
        "Homework",
        "Implement a small service, document its API, and submit the repository link.",
    )

    assert len(rubric.criteria) == 3
    assert all(criterion.max_points > 0 for criterion in rubric.criteria)
    assert rubric.total_points == 10.0


def test_extract_url_text_converts_github_blob_to_raw(monkeypatch) -> None:
    requested_urls: list[str] = []

    class Response:
        headers = {"content-type": "text/markdown; charset=utf-8"}
        content = b"# Assignment\nImplement the API and provide tests."
        text = content.decode()
        url = "https://raw.githubusercontent.com/example/course/main/README.md"

        def raise_for_status(self) -> None:
            return None

    class Client:
        def __init__(self, **kwargs) -> None:
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args) -> None:
            return None

        def get(self, url: str) -> Response:
            requested_urls.append(url)
            return Response()

    monkeypatch.setattr(parsers.httpx, "Client", Client)

    text = extract_url_text("https://github.com/example/course/blob/main/README.md")

    assert requested_urls == ["https://raw.githubusercontent.com/example/course/main/README.md"]
    assert "Implement the API" in text


def test_criteria_normalization_uses_fallback_for_incomplete_llm_response() -> None:
    fallback = [
        Criterion(name="Requirements", description="", max_points=4),
        Criterion(name="Quality", description="", max_points=6),
    ]

    criteria = _normalize_criteria(
        [{"name": "Only one", "description": "", "max_points": 1}],
        fallback,
    )

    assert [item["name"] for item in criteria] == ["Requirements", "Quality"]
